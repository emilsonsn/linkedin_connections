from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook
from openpyxl.worksheet.worksheet import Worksheet

from src.models import PersonInfo


class ConnectionWorkbook:
    HEADERS = ["data_hora", "nome", "descricao", "status", "perfil_linkedin"]

    def __init__(self, output_path: Path | str) -> None:
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook, self.sheet = self._open_workbook()

    def append(self, person: PersonInfo, status: str) -> None:
        self.sheet.append(
            [
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                person.name,
                person.description,
                status,
                person.profile_url,
            ]
        )
        self.workbook.save(self.output_path)

    def _open_workbook(self) -> tuple[OpenpyxlWorkbook, Worksheet]:
        if self.output_path.exists():
            workbook = load_workbook(self.output_path)
            sheet = workbook.active
            return workbook, sheet

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "conexoes"
        sheet.append(self.HEADERS)
        return workbook, sheet

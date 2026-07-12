from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.models import PersonInfo
from src.storage import ConnectionWorkbook


class ConnectionWorkbookTest(unittest.TestCase):
    def test_appends_to_the_same_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "connections.xlsx"
            workbook = ConnectionWorkbook(output_path)
            workbook.append(PersonInfo("Ana", "Desenvolvedora", "https://linkedin.com/in/ana"), "pendente_confirmado")
            workbook.append(PersonInfo("Bruno", "Designer", "https://linkedin.com/in/bruno"), "dry_run")

            saved_sheet = load_workbook(output_path).active
            self.assertEqual(saved_sheet.max_row, 3)
            self.assertEqual(saved_sheet["B2"].value, "Ana")
            self.assertEqual(saved_sheet["D3"].value, "dry_run")
